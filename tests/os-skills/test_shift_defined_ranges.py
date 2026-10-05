"""Shift direct defined-name references along with unpacked worksheets."""

import importlib.util
import subprocess
import sys
import tempfile
import unittest
import xml.dom.minidom
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

SCRIPT = (
    Path(__file__).resolve().parents[2] / "src/os-skills/others/xlsx/scripts/xlsx_shift_rows.py"
)
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


class DefinedRangeShiftTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        spec = importlib.util.spec_from_file_location("defined_shift", SCRIPT)
        cls.shift = importlib.util.module_from_spec(spec)
        previous = sys.path[:]
        try:
            sys.path.insert(0, str(SCRIPT.parent))
            spec.loader.exec_module(cls.shift)
        finally:
            sys.path[:] = previous

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / "xl").mkdir()
        self.path = self.root / "xl/workbook.xml"

    def names(self, values):
        workbook = ET.Element(f"{{{NS}}}workbook")
        names = ET.SubElement(workbook, f"{{{NS}}}definedNames")
        for name, value in values.items():
            node = ET.SubElement(
                names, f"{{{NS}}}definedName", name=name, localSheetId="0", hidden="1"
            )
            node.text = value
        ET.ElementTree(workbook).write(self.path, encoding="utf-8", xml_declaration=True)

    def values(self):
        return {
            x.get("name"): x.text
            for x in ET.parse(self.path).getroot().iter(f"{{{NS}}}definedName")
        }

    def run_shift(self, operation="insert", at=5, count=2):
        result = subprocess.run(
            [sys.executable, str(SCRIPT), str(self.root), operation, str(at), str(count)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=10,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return result

    def test_direct_cells_and_ranges_preserve_absolute_markers(self):
        self.names({"Region": "Sheet1!$A$3:$C$7", "Cell": "$D5", "Before": "A1:B4"})
        self.run_shift()
        self.assertEqual(
            self.values(), {"Region": "Sheet1!$A$3:$C$9", "Cell": "$D7", "Before": "A1:B4"}
        )
        for node in ET.parse(self.path).getroot().iter(f"{{{NS}}}definedName"):
            self.assertEqual(node.get("localSheetId"), "0")
            self.assertEqual(node.get("hidden"), "1")

    def test_print_titles_and_union_preserve_quoted_sheet_names(self):
        self.names(
            {
                "Print": "'Q!1 FY2025'!$1:$6, 'Q!1 FY2025'!$A:$C",
                "Escaped": "='O''Brien, FY2025'!$B$5,'O''Brien, FY2025'!D7",
            }
        )
        self.run_shift()
        self.assertEqual(self.values()["Print"], "'Q!1 FY2025'!$1:$8, 'Q!1 FY2025'!$A:$C")
        self.assertEqual(self.values()["Escaped"], "='O''Brien, FY2025'!$B$7,'O''Brien, FY2025'!D9")

    def test_deletion_uses_existing_shift_policy(self):
        self.names({"Rows": "Sheet1!$5:$9", "Union": "A2,A8:B10"})
        self.run_shift("delete", 5, 2)
        self.assertEqual(self.values(), {"Rows": "Sheet1!$3:$7", "Union": "A2,A6:B8"})

    def test_complex_external_3d_and_invalid_ranges_are_not_rewritten(self):
        values = {
            "Formula": "OFFSET(Sheet1!$A$5,0,0,5,1)",
            "Constant": '"A5"',
            "External": "'[Book.xlsx]Sheet1'!$A$5",
            "ThreeD": "Sheet1:Sheet3!A5",
            "Partial": "Sheet1!A5,OFFSET(A5,0,0)",
            "Grid": "XFE5:A9",
            "RowZero": "$0:$5",
            "Name": "Revenue",
            "Overflow": "A1048576",
        }
        self.names(values)
        original = self.path.read_bytes()
        self.run_shift()
        self.assertEqual(self.values(), values)
        self.assertEqual(self.path.read_bytes(), original)

    def test_shifting_preserves_unrelated_constant_whitespace_and_carriage_returns(self):
        self.names({"Range": "Sheet1!A5:B8", "Constant": '"first\n\n   \nlast"'})
        content = self.path.read_text(encoding="utf-8").replace("last", "last&#13;line")
        self.path.write_text(content, encoding="utf-8")
        original = self.values()["Constant"]
        self.run_shift()
        self.assertEqual(self.values()["Constant"], original)
        self.assertIn(b"&#13;", self.path.read_bytes())

    def test_shifting_preserves_compatibility_prefixes_comments_and_extensions(self):
        self.path.write_text(
            f'<?xml version="1.0" encoding="utf-8" standalone="yes"?>'
            f'<s:workbook xmlns:s="{NS}" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
            'xmlns:x15="urn:test:x15" xmlns:alias="urn:test:alias" mc:Ignorable="x15 alias">'
            '<!--keep this comment--><s:definedNames><s:definedName name="Range">'
            "Sheet1!A5:B8</s:definedName></s:definedNames>"
            '<s:extLst><s:ext uri="fixture"><alias:payload/></s:ext></s:extLst></s:workbook>',
            encoding="utf-8",
        )
        self.run_shift()
        document = xml.dom.minidom.parse(str(self.path))
        root = document.documentElement
        self.assertEqual(root.tagName, "s:workbook")
        self.assertEqual(root.getAttribute("xmlns:x15"), "urn:test:x15")
        self.assertEqual(root.getAttribute("xmlns:alias"), "urn:test:alias")
        self.assertEqual(root.getAttribute("mc:Ignorable"), "x15 alias")
        self.assertEqual(document.getElementsByTagName("alias:payload").length, 1)
        self.assertIn("<!--keep this comment-->", self.path.read_text(encoding="utf-8"))
        self.assertTrue(document.standalone)

    def test_preview_counts_changes_without_writing(self):
        self.names({"Range": "Sheet1!$A$5:$B$8", "Constant": "42"})
        original = self.path.read_bytes()
        self.assertEqual(self.shift.process_defined_names(str(self.path), 5, 2, dry_run=True), 1)
        self.assertEqual(self.path.read_bytes(), original)

    def test_whole_columns_and_unchanged_names_keep_file_bytes(self):
        self.names({"Columns": "Sheet1!$A:$XFD", "Before": "A1:B4"})
        original = self.path.read_bytes()
        self.run_shift()
        self.assertEqual(self.path.read_bytes(), original)

    def test_real_excel_print_area_titles_and_defined_ranges_track_cells(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Q!1 FY2025"
        sheet["A5"] = 10
        sheet["B7"] = 20
        sheet.print_area = "A3:C7"
        sheet.print_title_rows = "1:6"
        sheet.print_title_cols = "A:C"
        workbook.defined_names.add(DefinedName("Region", attr_text="'Q!1 FY2025'!$A$5:$B$7"))
        original = self.root / "input.xlsx"
        workbook.save(original)
        workbook.close()
        with zipfile.ZipFile(original) as archive:
            archive.extractall(self.root)
        self.run_shift()
        output = self.root / "output.xlsx"
        with zipfile.ZipFile(output, "w") as archive:
            for path in self.root.rglob("*"):
                if path.is_file() and path.suffix != ".xlsx":
                    archive.write(path, path.relative_to(self.root).as_posix())
        loaded = load_workbook(output)
        self.addCleanup(loaded.close)
        sheet = loaded["Q!1 FY2025"]
        self.assertEqual(sheet["A7"].value, 10)
        self.assertEqual(sheet["B9"].value, 20)
        self.assertEqual(loaded.defined_names["Region"].attr_text, "'Q!1 FY2025'!$A$7:$B$9")
        self.assertEqual(sheet.print_title_rows, "$1:$8")
        self.assertEqual(sheet.print_title_cols, "$A:$C")
        self.assertIn("$A$3:$C$9", str(sheet.print_area))


if __name__ == "__main__":
    unittest.main()
